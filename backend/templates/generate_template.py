"""Membuat template BA Drop. Jalankan dari root backend dengan Python 3."""

from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_ROW_HEIGHT_RULE, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt

OUTPUT = Path(__file__).with_name("ba-drop-template.docx")
ASSET_DIR = Path(__file__).with_name("assets")


def set_cell_border(cell, **edges):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_borders = tc_pr.first_child_found_in("w:tcBorders")
    if tc_borders is None:
        tc_borders = OxmlElement("w:tcBorders")
        tc_pr.append(tc_borders)
    for edge, settings in edges.items():
        tag = f"w:{edge}"
        element = tc_borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            tc_borders.append(element)
        for key, value in settings.items():
            element.set(qn(f"w:{key}"), str(value))


def set_paragraph_bottom_border(paragraph, size=10):
    p_pr = paragraph._p.get_or_add_pPr()
    borders = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), str(size))
    bottom.set(qn("w:space"), "5")
    bottom.set(qn("w:color"), "000000")
    borders.append(bottom)
    p_pr.append(borders)


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shading = OxmlElement("w:shd")
    shading.set(qn("w:fill"), fill)
    tc_pr.append(shading)


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    header = OxmlElement("w:tblHeader")
    header.set(qn("w:val"), "true")
    tr_pr.append(header)


def set_cell_text(cell, text, *, bold=False, size=8, align=WD_ALIGN_PARAGRAPH.LEFT):
    cell.text = ""
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    paragraph = cell.paragraphs[0]
    paragraph.alignment = align
    paragraph.paragraph_format.space_after = Pt(0)
    paragraph.paragraph_format.space_before = Pt(0)
    run = paragraph.add_run(text)
    run.bold = bold
    run.font.name = "Arial"
    run.font.size = Pt(size)
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "Arial")


def remove_table_borders(table):
    for row in table.rows:
        for cell in row.cells:
            none = {"val": "nil"}
            set_cell_border(cell, top=none, left=none, bottom=none, right=none, insideH=none, insideV=none)


def set_cell_margins(cell, top=80, start=80, bottom=80, end=80):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_keep_with_next(paragraph, value=True):
    p_pr = paragraph._p.get_or_add_pPr()
    keep = p_pr.find(qn("w:keepNext"))
    if keep is None:
        keep = OxmlElement("w:keepNext")
        p_pr.append(keep)
    keep.set(qn("w:val"), "true" if value else "false")


def set_run_font(run, size=9, bold=False):
    run.bold = bold
    run.font.name = "Arial"
    run.font.size = Pt(size)
    fonts = run._element.get_or_add_rPr().rFonts
    fonts.set(qn("w:ascii"), "Arial")
    fonts.set(qn("w:hAnsi"), "Arial")
    fonts.set(qn("w:eastAsia"), "Arial")


def add_signature_cell(cell, organization, role, name, nik):
    cell.text = ""
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
    set_cell_margins(cell, top=40, start=60, bottom=40, end=60)
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(0)
    for line in organization:
        run = p.add_run(line + "\n")
        set_run_font(run, size=9, bold=True)
    for line in role:
        run = p.add_run(line + "\n")
        set_run_font(run, size=9, bold=True)
    spacer = p.add_run("\n\n\n")
    set_run_font(spacer, size=9)
    name_run = p.add_run(name + "\n")
    set_run_font(name_run, size=9, bold=True)
    name_run.underline = True
    nik_run = p.add_run(nik)
    set_run_font(nik_run, size=9, bold=True)


doc = Document()
section = doc.sections[0]
section.page_width = Cm(21)
section.page_height = Cm(29.7)
section.top_margin = Cm(2.35)
section.bottom_margin = Cm(1.2)
section.left_margin = Cm(1.1)
section.right_margin = Cm(1.1)
section.header_distance = Cm(0.45)
section.different_first_page_header_footer = False

normal = doc.styles["Normal"]
normal.font.name = "Arial"
normal.font.size = Pt(9)
normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Arial")

# Header Word digunakan supaya kedua logo otomatis berulang pada setiap
# halaman, termasuk halaman tambahan akibat tabel yang memanjang.
header = section.header
header.is_linked_to_previous = False
header_paragraph = header.paragraphs[0]
header_paragraph.paragraph_format.space_after = Pt(0)
header_table = header.add_table(rows=1, cols=2, width=Cm(18.8))
header_table.alignment = WD_TABLE_ALIGNMENT.CENTER
header_table.autofit = False
header_table.columns[0].width = Cm(9.4)
header_table.columns[1].width = Cm(9.4)
remove_table_borders(header_table)
for index, (image_name, width, align) in enumerate([
    ("infranexia.png", Cm(4.1), WD_ALIGN_PARAGRAPH.LEFT),
    ("telkomakses.png", Cm(3.8), WD_ALIGN_PARAGRAPH.RIGHT),
]):
    cell = header_table.cell(0, index)
    cell.width = Cm(9.4)
    cell.text = ""
    set_cell_margins(cell, top=0, start=0, bottom=0, end=0)
    paragraph = cell.paragraphs[0]
    paragraph.alignment = align
    paragraph.paragraph_format.space_after = Pt(0)
    paragraph.add_run().add_picture(str(ASSET_DIR / image_name), width=width)

title = doc.add_paragraph()
title.alignment = WD_ALIGN_PARAGRAPH.CENTER
title.paragraph_format.space_before = Pt(3)
title.paragraph_format.space_after = Pt(8)
title_run = title.add_run("BERITA ACARA DROP LOP PEKERJAAN")
title_run.bold = True
title_run.font.name = "Arial"
title_run.font.size = Pt(14)
set_paragraph_bottom_border(title)

info = doc.add_table(rows=4, cols=3)
info.alignment = WD_TABLE_ALIGNMENT.CENTER
info.autofit = False
widths = [Cm(2.8), Cm(0.45), Cm(15.0)]
labels = [
    ("Proyek", "{judulProyek}"),
    ("No. Kontrak", "{nomorKontrak}"),
    ("No. SP", "{nomorSp}"),
    ("Pelaksana", "{pelaksana}"),
]
for row, (label, value) in zip(info.rows, labels):
    for cell, width in zip(row.cells, widths):
        cell.width = width
    set_cell_text(row.cells[0], label, bold=True, size=9)
    set_cell_text(row.cells[1], ":", bold=True, size=9, align=WD_ALIGN_PARAGRAPH.CENTER)
    set_cell_text(row.cells[2], value, bold=True, size=9)
remove_table_borders(info)

rule = doc.add_paragraph()
rule.paragraph_format.space_before = Pt(4)
rule.paragraph_format.space_after = Pt(9)
set_paragraph_bottom_border(rule)

opening = doc.add_paragraph()
opening.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
opening.paragraph_format.first_line_indent = Cm(1.5)
opening.paragraph_format.line_spacing = 1.15
opening.paragraph_format.space_after = Pt(8)
opening.add_run("Pada hari ini ")
opening.add_run("{tanggalBeritaAcara}").bold = True
opening.add_run(", Telkomsel, TIF dan Telkom Akses sepakat bahwa Project ")
opening.add_run("{judulProyek}").bold = True
opening.add_run(", District ")
opening.add_run("{district}").bold = True
opening.add_run(" dengan lokasi sebagai berikut:")

headers = ["NO", "WOK", "Tipe\nDesain", "Nama Proyek", "Keterangan Drop", "IHLD Lop\nID", "Jml\nODP", "Jml\nPort", "Total BOQ"]
widths = [Cm(0.95), Cm(1.2), Cm(1.45), Cm(3.2), Cm(3.45), Cm(2.0), Cm(1.05), Cm(1.05), Cm(2.35)]
table = doc.add_table(rows=3, cols=9)
table.alignment = WD_TABLE_ALIGNMENT.CENTER
table.autofit = False
table.style = "Table Grid"

for cell, text, width in zip(table.rows[0].cells, headers, widths):
    cell.width = width
    set_cell_text(cell, text, bold=True, size=7.5, align=WD_ALIGN_PARAGRAPH.CENTER)
    set_cell_shading(cell, "F2F2F2")
set_repeat_table_header(table.rows[0])

loop_values = ["{#items}{no}", "{wok}", "{tipeDesain}", "{namaProyek}", "{keteranganDrop}", "{ihldLopId}", "{jmlOdp}", "{jmlPort}", "{totalBoq}{/items}"]
for index, (cell, text, width) in enumerate(zip(table.rows[1].cells, loop_values, widths)):
    cell.width = width
    set_cell_text(cell, text, size=7.5, align=WD_ALIGN_PARAGRAPH.CENTER if index != 4 else WD_ALIGN_PARAGRAPH.LEFT)

total_values = ["", "", "", "", "", "TOTAL", "{totalOdp}", "{totalPort}", "{totalBoq}"]
for cell, text, width in zip(table.rows[2].cells, total_values, widths):
    cell.width = width
    set_cell_text(cell, text, bold=True, size=7.5, align=WD_ALIGN_PARAGRAPH.CENTER)

closing = doc.add_paragraph()
closing.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
closing.paragraph_format.space_before = Pt(9)
closing.paragraph_format.line_spacing = 1.15
closing.add_run("Demikian, untuk ")
closing.add_run("{jumlahLop}").bold = True
closing.add_run(" LOP tersebut tidak dapat dilanjutkan pekerjaannya berdasarkan kendala keterangan di atas.")

# Halaman tanda tangan dibuat terpisah agar posisi blok tetap konsisten walaupun
# tabel proyek memanjang ke beberapa halaman.
doc.add_page_break()
conclusion = doc.add_paragraph()
conclusion.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
conclusion.paragraph_format.line_spacing = 1.15
conclusion.paragraph_format.space_after = Pt(26)
conclusion.add_run("Setelah berkoordinasi dengan Telkomsel dan TIF pada project yang terdapat dalam ")
conclusion.add_run("{judulProyek}").bold = True
conclusion.add_run(", untuk ")
conclusion.add_run("{jumlahLop} LOP").bold = True
conclusion.add_run(" tersebut tidak dapat dilanjutkan pekerjaannya berdasarkan kendala keterangan di atas. ")
conclusion.add_run("Demikian Berita Acara Drop Pekerjaan ini dibuat sesuai keadaan sebenarnya dan untuk dipergunakan sebagaimana mestinya.")

date_line = doc.add_paragraph()
date_line.alignment = WD_ALIGN_PARAGRAPH.RIGHT
date_line.paragraph_format.space_after = Pt(26)
date_run = date_line.add_run("{tanggalDokumen}")
set_run_font(date_run, size=9)

top_signatures = doc.add_table(rows=1, cols=2)
top_signatures.alignment = WD_TABLE_ALIGNMENT.CENTER
top_signatures.autofit = False
for cell in top_signatures.rows[0].cells:
    cell.width = Cm(9.2)
remove_table_borders(top_signatures)
add_signature_cell(
    top_signatures.cell(0, 0),
    ["PT. TELKOM AKSES"],
    ["SO PROJECT DEPLOYMENT", "AREA JAWA BALI"],
    "{signer1Name}",
    "{signer1Nik}",
)
add_signature_cell(
    top_signatures.cell(0, 1),
    ["PT. TELKOM INFRASTRUKTUR", "INDONESIA"],
    ["OFF 2 OM & PROJECT OUTSIDE PLANT"],
    "{signer2Name}",
    "{signer2Nik}",
)

approval = doc.add_paragraph()
approval.alignment = WD_ALIGN_PARAGRAPH.CENTER
approval.paragraph_format.space_before = Pt(14)
approval.paragraph_format.space_after = Pt(5)
approval_run = approval.add_run("Mengetahui dan Menyetujui,")
set_run_font(approval_run, size=9)

bottom_signatures = doc.add_table(rows=1, cols=3)
bottom_signatures.alignment = WD_TABLE_ALIGNMENT.CENTER
bottom_signatures.autofit = False
for cell in bottom_signatures.rows[0].cells:
    cell.width = Cm(6.1)
remove_table_borders(bottom_signatures)
add_signature_cell(
    bottom_signatures.cell(0, 0),
    ["PT. TELKOM AKSES"],
    ["MGR PROJECT DEPLOYMENT", "AREA JAWA BALI"],
    "{signer3Name}",
    "{signer3Nik}",
)
add_signature_cell(
    bottom_signatures.cell(0, 1),
    ["PT. TELKOM INFRASTRUKTUR", "INDONESIA"],
    ["MGR NEW ACCESS NW DESIGN", "ENGINEERING"],
    "{signer4Name}",
    "{signer4Nik}",
)
add_signature_cell(
    bottom_signatures.cell(0, 2),
    ["PT. TELKOMSEL"],
    ["MANAGER NETWORK", "BROADBAND ACCESS", "DEPLOYMENT"],
    "{signer5Name}",
    "{signer5Nik}",
)

# Lampiran evidence: dua kotak per baris seperti dokumen acuan. Caption setiap
# kotak dibentuk dari Nama Proyek dan Keterangan Drop pada data terpilih.
doc.add_page_break()
appendix_title = doc.add_paragraph()
appendix_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
appendix_title.paragraph_format.space_after = Pt(10)
appendix_run = appendix_title.add_run("LAMPIRAN EVIDENCE")
set_run_font(appendix_run, size=10, bold=True)
appendix_run.underline = True
set_keep_with_next(appendix_title)

evidence_table = doc.add_table(rows=1, cols=2)
evidence_table.alignment = WD_TABLE_ALIGNMENT.CENTER
evidence_table.autofit = False
evidence_table.style = "Table Grid"
evidence_row = evidence_table.rows[0]
evidence_row.height = Cm(7.5)
evidence_row.height_rule = WD_ROW_HEIGHT_RULE.AT_LEAST
for cell, text in zip(evidence_row.cells, ["{#evidenceRows}{leftCaption}", "{rightCaption}{/evidenceRows}"]):
    cell.width = Cm(9.15)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.BOTTOM
    set_cell_margins(cell, top=100, start=100, bottom=100, end=100)
    set_cell_text(cell, text, size=9, align=WD_ALIGN_PARAGRAPH.LEFT)

doc.save(OUTPUT)
print(f"Template dibuat: {OUTPUT}")
